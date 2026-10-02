"""Trusted format conversion worker. Resource limits precede parser imports."""
import ctypes
import io
import json
import sys
import zipfile
from pathlib import Path

MEMORY = 256 * 1024 * 1024
MAX_CHARS = 500_000
_job_handle = None


class LimitError(Exception):
    pass


def set_limits():
    global _job_handle
    if sys.platform == "win32":
        from ctypes import wintypes as w
        size = ctypes.c_size_t
        class Basic(ctypes.Structure):
            _fields_ = [("process_time", ctypes.c_longlong), ("job_time", ctypes.c_longlong),
                        ("flags", w.DWORD), ("min_ws", size), ("max_ws", size),
                        ("active", w.DWORD), ("affinity", size), ("priority", w.DWORD), ("scheduling", w.DWORD)]
        class IO(ctypes.Structure):
            _fields_ = [(name, ctypes.c_ulonglong) for name in ("read", "write", "other", "read_bytes", "write_bytes", "other_bytes")]
        class Extended(ctypes.Structure):
            _fields_ = [("basic", Basic), ("io", IO), ("process_memory", size), ("job_memory", size),
                        ("peak_process", size), ("peak_job", size)]
        kernel = ctypes.WinDLL("kernel32", use_last_error=True)
        kernel.CreateJobObjectW.argtypes = [ctypes.c_void_p, w.LPCWSTR]
        kernel.CreateJobObjectW.restype = w.HANDLE
        kernel.SetInformationJobObject.argtypes = [w.HANDLE, ctypes.c_int, ctypes.c_void_p, w.DWORD]
        kernel.SetInformationJobObject.restype = w.BOOL
        kernel.GetCurrentProcess.restype = w.HANDLE
        kernel.AssignProcessToJobObject.argtypes = [w.HANDLE, w.HANDLE]
        kernel.AssignProcessToJobObject.restype = w.BOOL
        info = Extended()
        info.basic.flags = 0x2000 | 0x100 | 0x2 | 0x8  # kill-on-close, memory, CPU, one process
        info.basic.process_time = 10 * 10_000_000
        info.basic.active = 1
        info.process_memory = MEMORY
        handle = kernel.CreateJobObjectW(None, None)
        if not handle or not kernel.SetInformationJobObject(handle, 9, ctypes.byref(info), ctypes.sizeof(info)):
            raise OSError(ctypes.get_last_error(), "无法建立 Windows 资源限制")
        if not kernel.AssignProcessToJobObject(handle, kernel.GetCurrentProcess()):
            raise OSError(ctypes.get_last_error(), "无法加入 Windows Job Object")
        _job_handle = handle
    elif sys.platform.startswith("linux"):
        import resource
        resource.setrlimit(resource.RLIMIT_AS, (MEMORY, MEMORY))
        resource.setrlimit(resource.RLIMIT_CPU, (10, 11))
        resource.setrlimit(resource.RLIMIT_CORE, (0, 0))
    else:
        raise OSError("当前系统不支持文件转换的硬性资源限制")


def text(value):
    if value is None:
        return ""
    return str(value)


def table(rows, fmt):
    if not rows:
        return ""
    width = max(map(len, rows))
    if fmt == "txt":
        return "\n".join("\t".join(text(c).replace("\t", " ").replace("\n", " ") for c in row) for row in rows)
    def line(row):
        cells = [text(c).replace("\\", "\\\\").replace("|", "\\|").replace("\r", "").replace("\n", "<br>") for c in row]
        return "| " + " | ".join(cells + [""] * (width - len(cells))) + " |"
    return "\n".join([line(rows[0]), line(["---"] * width), *(line(row) for row in rows[1:])])


def identify(path):
    with open(path, "rb") as source:
        signature = source.read(8)
    if signature.startswith(b"%PDF-"):
        return "pdf"
    if not zipfile.is_zipfile(path):
        raise ValueError("只支持 PDF、DOCX 或 XLSX；不支持旧版或加密文件")
    with zipfile.ZipFile(path) as archive:
        entries = archive.infolist()
        if len(entries) > 2000 or sum(item.file_size for item in entries) > 50 * 1024 * 1024:
            raise LimitError("文档解压大小或条目数量超过限制")
        if any(item.flag_bits & 1 for item in entries):
            raise ValueError("不支持加密文档")
        names = {item.filename for item in entries}
        if any("vbaproject" in name.lower() for name in names):
            raise ValueError("不支持宏文件")
        if "[Content_Types].xml" not in names:
            raise ValueError("无效的 Office 文档")
        if "word/document.xml" in names and "xl/workbook.xml" not in names:
            kind = "docx"
        elif "xl/workbook.xml" in names and "word/document.xml" not in names:
            kind = "xlsx"
        else:
            raise ValueError("无法识别 Office 文档类型")
        types = archive.read("[Content_Types].xml").lower()
        if b"macroenabled" in types or b"template.main+xml" in types:
            raise ValueError("不支持宏或模板文件")
        return kind


def convert(path, fmt, filename):
    if fmt not in {"markdown", "txt"}:
        raise ValueError("format 必须为 markdown 或 txt")
    kind = identify(path)
    suffix = Path(filename).suffix.lower()
    if suffix and suffix not in {".pdf", ".docx", ".xlsx"}:
        raise ValueError("不支持的文件扩展名")
    if suffix and suffix != "." + kind:
        raise ValueError("文件名与实际文档类型不匹配")
    parts, warnings = [], []
    stats = {"pages": 0, "worksheets": 0, "cells": 0, "characters": 0}
    def append(value):
        stats["characters"] += len(value) + (2 if parts else 0)
        if stats["characters"] > MAX_CHARS:
            raise LimitError("转换文本超过 500,000 字符")
        parts.append(value)
    if kind == "pdf":
        from pypdf import PdfReader
        reader = PdfReader(path)
        if reader.is_encrypted:
            raise ValueError("不支持加密 PDF")
        if len(reader.pages) > 100:
            raise LimitError("PDF 不能超过 100 页")
        found = False
        stats["pages"] = len(reader.pages)
        for index, page in enumerate(reader.pages, 1):
            content = (page.extract_text() or "").strip()
            found |= bool(content)
            if not content:
                warnings.append(f"第 {index} 页没有可提取文字，可能是扫描页；未执行 OCR")
            append((f"## 第 {index} 页\n\n" if fmt == "markdown" else f"第 {index} 页\n") + content)
        if not found:
            raise ValueError("PDF 没有可提取文字，扫描件需要 OCR，当前不支持")
    elif kind == "docx":
        from docx import Document
        from docx.text.paragraph import Paragraph
        document = Document(path)
        for item in document.iter_inner_content():
            if isinstance(item, Paragraph):
                value = item.text
                style = item.style.name if item.style else ""
                if fmt == "markdown" and style.startswith("Heading ") and style[8:].isdigit():
                    value = "#" * min(int(style[8:]), 6) + " " + value
                if value:
                    append(value)
            else:
                append(table([[cell.text for cell in row.cells] for row in item.rows], fmt))
        warnings.append("仅提取正文段落、标题和表格；不还原图片、页眉页脚或复杂排版")
    else:
        from openpyxl import load_workbook
        from openpyxl.utils import get_column_letter
        data = Path(path).read_bytes()
        values = load_workbook(io.BytesIO(data), read_only=True, data_only=True, keep_links=False)
        formulas = None
        try:
            formulas = load_workbook(io.BytesIO(data), read_only=True, data_only=False, keep_links=False)
            if len(values.worksheets) > 20:
                raise LimitError("XLSX 不能超过 20 个工作表")
            stats["worksheets"] = len(values.worksheets)
            missing_formula = False
            for sheet, formula_sheet in zip(values.worksheets, formulas.worksheets):
                sheet.reset_dimensions()
                formula_sheet.reset_dimensions()
                rows, chars = [], 0
                for row, original in zip(sheet.iter_rows(), formula_sheet.iter_rows()):
                    stats["cells"] += len(row)
                    if stats["cells"] > 50_000:
                        raise LimitError("XLSX 累计遍历单元格超过 50,000")
                    converted = []
                    for cell, formula in zip(row, original):
                        if formula.data_type == "f" and cell.value is None:
                            missing_formula = True
                        value = text(cell.value)
                        chars += len(value) + 8
                        if chars > MAX_CHARS:
                            raise LimitError("工作表转换文本超过限制")
                        converted.append(value)
                    rows.append(converted)
                if fmt == "markdown" and rows:
                    width = max(map(len, rows))
                    rows.insert(0, [get_column_letter(i + 1) for i in range(width)])
                append(("## " if fmt == "markdown" else "") + sheet.title + "\n\n" + table(rows, fmt))
            if missing_formula:
                warnings.append("部分公式没有已保存的计算结果，已输出为空；仅显示已保存的结果，不计算公式")
        finally:
            values.close()
            if formulas is not None:
                formulas.close()
    result = "\n\n".join(parts)
    stats["characters"] = len(result)
    stem = Path(filename).stem or "document"
    return {"result": result, "format": fmt, "filename": stem + (".md" if fmt == "markdown" else ".txt"),
            "source_type": kind, "stats": stats, "warnings": warnings}


def main():
    try:
        set_limits()
    except Exception:
        print(json.dumps({"error": {"status": 503, "code": "DOCUMENT_UNAVAILABLE", "message": "无法建立文件转换资源限制"}}))
        return
    try:
        request = json.loads(sys.stdin.buffer.read(16 * 1024))
        if request.get("file_url"):
            from app.image_source import fetch_file, ImageTooLargeError
            try:
                downloaded = fetch_file(request["file_url"], 5 * 1024 * 1024)
            except ImageTooLargeError:
                sys.stdout.buffer.write(json.dumps({"error": {"status": 413, "code": "FILE_TOO_LARGE", "message": "文件不能超过 5 MiB"}}).encode())
                return
            if not downloaded:
                raise ValueError("文件不能为空")
            Path(request["path"]).write_bytes(downloaded)
            del downloaded
        sys.stdout.buffer.write(b'{"stage":"parsing"}\n')
        sys.stdout.buffer.flush()
        output = convert(request["path"], request["format"], request["filename"])
    except (LimitError, MemoryError):
        output = {"error": {"status": 413, "code": "DOCUMENT_LIMIT_EXCEEDED", "message": "文档处理超过容量或内存限制"}}
    except Exception as exc:
        output = {"error": {"status": 400, "code": "INVALID_INPUT", "message": f"无法解析文档：{str(exc)[:300]}"}}
    sys.stdout.buffer.write(json.dumps(output, ensure_ascii=False).encode("utf-8"))


if __name__ == "__main__":
    main()
