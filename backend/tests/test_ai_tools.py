import asyncio
import io
import json
import time
import subprocess
import sys
import zipfile
from datetime import datetime, timezone
from decimal import Decimal

import httpx
import pytest
from docx import Document
from fastapi.testclient import TestClient
from openpyxl import Workbook
from pypdf import PdfWriter
from pypdf.generic import DictionaryObject, NameObject, DecodedStreamObject, ArrayObject, NumberObject, TextStringObject

from app import current_time, currency, documents, image_source
from app.document_worker import convert, LimitError
from app.main import app
from app.schemas import CurrencyInput
from app.tool_errors import ToolFailure

client = TestClient(app)


@pytest.mark.parametrize("instant,expected,weekday", [
    ("2026-10-01T16:00:00+00:00", "2026-10-02T00:00:00.000+08:00", 5),
    ("2026-12-31T16:00:00.123000+00:00", "2027-01-01T00:00:00.123+08:00", 5),
])
def test_time(instant, expected, weekday, monkeypatch):
    value = datetime.fromisoformat(instant)
    monkeypatch.setattr(current_time, "now", lambda: current_time.describe(value))
    response = client.get("/api/v1/time/now")
    data = response.json()
    assert response.headers["cache-control"] == "no-store"
    assert data["iso_datetime"] == expected and data["weekday"] == weekday
    assert int(data["unix_milliseconds"]) // 1000 == int(data["unix_seconds"])
    assert data["date"] == expected[:10] and data["time"] == expected[11:23]


def fake_service():
    service = currency.CurrencyService()
    calls = []
    async def fetch(path):
        calls.append(path)
        await asyncio.sleep(0)
        if path == "/currencies":
            return [{"iso_code": "CNY", "name": "Yuan"}, {"iso_code": "USD", "name": "Dollar"}]
        return {"base": "CNY", "quote": "USD", "rate": Decimal("0.145"), "date": "2026-10-01"}
    service._fetch = fetch
    return service, calls


def test_currency_precision_cache_and_coalescing():
    async def scenario():
        service, calls = fake_service()
        body = CurrencyInput(amount="1", from_currency="cny", to_currency="usd")
        results = await asyncio.gather(*(service.convert(body) for _ in range(8)))
        assert all(result["result"] == "0.15" for result in results)
        assert calls == ["/currencies", "/rate/cny/usd"]
        assert (await service.convert(body))["cached"] is True
        same = await service.convert(CurrencyInput(amount="999999999999999.12345678", from_currency="USD", to_currency="USD", precision=8))
        assert same["result"] == "999999999999999.12345678" and same["rate_date"] is None
        assert same["rate"] == "1" and same["source"] == "identity"
        with pytest.raises(ValueError):
            await service.convert(CurrencyInput(amount="1", from_currency="ABC", to_currency="USD"))
    asyncio.run(scenario())


def test_currency_stale_failure_and_capacity():
    async def scenario():
        service, _ = fake_service()
        await service.currencies()
        await service.rate("CNY", "USD")
        key = ("CNY", "USD")
        _, entry = service.rates[key]
        async def fail(path):
            raise httpx.ConnectError("offline")
        service._fetch = fail
        service.rates[key] = (time.monotonic() - 7200, entry)
        service.directory = (time.monotonic() - 90000, service.directory[1])
        response = await service.convert(CurrencyInput(amount="1", from_currency="CNY", to_currency="USD"))
        assert response["cached"] and response["stale"]
        service.rates[key] = (time.monotonic() - 90000, entry)
        with pytest.raises(ToolFailure) as error:
            await service.rate(*key)
        assert error.value.code == "EXCHANGE_RATE_UNAVAILABLE"
        service, _ = fake_service()
        for index in range(300):
            service.rates[(str(index), "USD")] = (0, {})
        await service.rate("CNY", "USD")
        assert len(service.rates) == 256
    asyncio.run(scenario())


@pytest.mark.parametrize("updates", [
    {"amount": "-1"}, {"amount": "NaN"}, {"amount": "1e2"}, {"amount": "1.123456789"},
    {"amount": "1000000000000000"}, {"amount": 1}, {"precision": 9}, {"precision": "2"}, {"precision": True},
    {"from_currency": "US"},
])
def test_currency_validation(updates):
    body = {"amount": "1", "from_currency": "CNY", "to_currency": "USD", **updates}
    result = client.post("/api/v1/currency/convert", json=body)
    assert result.status_code == 422 and result.json()["code"] == "VALIDATION_ERROR"


def test_currency_http(monkeypatch):
    service, _ = fake_service()
    monkeypatch.setattr(currency, "service", service)
    with TestClient(app) as test:
        assert test.get("/api/v1/currency/currencies").json()["currencies"][0]["code"] == "CNY"
        data = test.post("/api/v1/currency/convert", json={"amount": "1", "from_currency": "CNY", "to_currency": "USD"}).json()
        assert data["result"] == "0.15" and data["rate_date"] == "2026-10-01"


def make_docx():
    doc = Document()
    doc.add_heading("中文报告", level=1)
    doc.add_paragraph("前段")
    cells = doc.add_table(rows=2, cols=2).rows
    cells[0].cells[0].text = "字段"
    cells[0].cells[1].text = "值"
    cells[1].cells[0].text = "管道|换行"
    cells[1].cells[1].text = "你好\n世界"
    doc.add_paragraph("后段")
    output = io.BytesIO(); doc.save(output)
    return output.getvalue()


def make_pdf(pages=1, encrypted=False):
    writer = PdfWriter()
    for index in range(pages):
        page = writer.add_blank_page(width=200, height=200)
        if index == 0:
            # Minimal CJK text fixture with explicit Unicode mapping.
            cmap = DecodedStreamObject()
            cmap.set_data(b"/CIDInit /ProcSet findresource begin 12 dict begin begincmap /CIDSystemInfo << /Registry (Adobe) /Ordering (UCS) /Supplement 0 >> def /CMapName /Test def /CMapType 2 def 1 begincodespacerange <0000> <FFFF> endcodespacerange 2 beginbfchar <0001> <4F60> <0002> <597D> endbfchar endcmap CMapName currentdict /CMap defineresource pop end end")
            font = DictionaryObject({NameObject("/Type"): NameObject("/Font"), NameObject("/Subtype"): NameObject("/Type0"), NameObject("/BaseFont"): NameObject("/Test"), NameObject("/Encoding"): NameObject("/Identity-H"), NameObject("/ToUnicode"): writer._add_object(cmap)})
            descendant = DictionaryObject({NameObject("/Type"): NameObject("/Font"), NameObject("/Subtype"): NameObject("/CIDFontType2"), NameObject("/BaseFont"): NameObject("/Test"), NameObject("/DW"): NumberObject(1000), NameObject("/CIDSystemInfo"): DictionaryObject({NameObject("/Registry"): TextStringObject("Adobe"), NameObject("/Ordering"): TextStringObject("Identity"), NameObject("/Supplement"): NumberObject(0)})})
            font[NameObject("/DescendantFonts")] = ArrayObject([writer._add_object(descendant)])
            page[NameObject("/Resources")] = DictionaryObject({NameObject("/Font"): DictionaryObject({NameObject("/F1"): writer._add_object(font)})})
            stream = DecodedStreamObject(); stream.set_data(b"BT /F1 12 Tf 10 100 Td <00010002> Tj ET")
            page[NameObject("/Contents")] = writer._add_object(stream)
    if encrypted:
        writer.encrypt("secret")
    output = io.BytesIO(); writer.write(output)
    return output.getvalue()


def test_documents_actual_worker_docx_and_xlsx():
    response = client.post("/api/v1/documents/convert", files={"file": ("报告.docx", make_docx())})
    assert response.status_code == 200, response.text
    data = response.json()
    assert data["filename"] == "报告.md" and data["result"].startswith("# 中文报告")
    assert data["result"].index("前段") < data["result"].index("字段") < data["result"].index("后段")
    assert "管道\\|换行" in data["result"] and "你好<br>世界" in data["result"]
    workbook = Workbook()
    workbook.active.append(["中文", "a|b", "=1+1"])
    workbook.create_sheet("第二页").append([123, "你好"])
    output = io.BytesIO(); workbook.save(output)
    response = client.post("/api/v1/documents/convert", files={"file": ("data.xlsx", output.getvalue())})
    assert response.status_code == 200, response.text
    data = response.json()
    assert "| A | B | C |" in data["result"] and "a\\|b" in data["result"]
    assert data["stats"]["worksheets"] == 2 and data["warnings"]
    plain = client.post("/api/v1/documents/convert", files={"file": ("report.docx", make_docx())}, data={"format": "txt"})
    assert "你好 世界" in plain.json()["result"] and plain.json()["filename"] == "report.txt"


def test_pdf_text_scanned_mixed_and_encrypted():
    data = client.post("/api/v1/documents/convert", files={"file": ("report.pdf", make_pdf(2))}).json()
    assert "你好" in data["result"] and "第 2 页" in data["warnings"][0]
    writer = PdfWriter(); writer.add_blank_page(width=200, height=200)
    scanned = io.BytesIO(); writer.write(scanned)
    for content in (scanned.getvalue(), make_pdf(encrypted=True)):
        response = client.post("/api/v1/documents/convert", files={"file": ("report.pdf", content)})
        assert response.status_code == 400 and response.json()["code"] == "INVALID_INPUT"


def test_document_sources_sizes_and_failure_recovery(monkeypatch):
    for params in [{}, {"data": {"file_url": "https://example.com/a.docx"}, "files": {"file": ("a.docx", make_docx())}},
                   {"files": {"file": ("a.pdf", b"bad")}}, {"files": {"file": ("a.xlsx", make_docx())}}]:
        assert client.post("/api/v1/documents/convert", **params).status_code == 400
    assert client.post("/api/v1/documents/convert", files={"file": ("a.pdf", b"x" * (5 * 1024 * 1024 + 1))}).status_code == 413
    assert client.post("/api/v1/documents/convert", content=b"x" * (6 * 1024 * 1024 + 1)).json()["code"] == "FILE_TOO_LARGE"
    assert client.post("/api/v1/documents/convert", files={"file": ("a.docx", make_docx())}, data={"format": "html"}).status_code == 422
    actual_worker = documents.run_worker
    async def fake_download(path, fmt, filename, request, file_url):
        assert file_url == "https://example.com/download?token=123"
        path.write_bytes(make_docx())
        return await actual_worker(path, fmt, filename, request)
    monkeypatch.setattr(documents, "run_worker", fake_download)
    response = client.post("/api/v1/documents/convert", data={"file_url": "https://example.com/download?token=123"})
    assert response.status_code == 200 and response.json()["source_type"] == "docx"
    assert client.get("/api/v1/health").status_code == 200


def test_document_capacity_limits(tmp_path, monkeypatch):
    path = tmp_path / "input"
    path.write_bytes(make_pdf(101))
    with pytest.raises(LimitError):
        convert(path, "markdown", "a.pdf")
    wb = Workbook()
    for _ in range(20): wb.create_sheet()
    wb.save(path)
    with pytest.raises(LimitError): convert(path, "txt", "a.xlsx")
    wb = Workbook()
    for _ in range(501): wb.active.append(["large"] * 100)
    wb.save(path)
    with pytest.raises(LimitError): convert(path, "txt", "a.xlsx")
    with zipfile.ZipFile(path, "w") as archive:
        for n in range(2001): archive.writestr(str(n), "")
    with pytest.raises(LimitError): convert(path, "txt", "a.docx")
    with zipfile.ZipFile(path, "w", compression=zipfile.ZIP_DEFLATED) as archive:
        archive.writestr("large", b"x" * (50 * 1024 * 1024 + 1))
    with pytest.raises(LimitError): convert(path, "txt", "a.docx")
    path.write_bytes(make_docx())
    import app.document_worker as worker
    monkeypatch.setattr(worker, "MAX_CHARS", 10)
    with pytest.raises(LimitError): convert(path, "txt", "a.docx")


def test_document_admission_before_body_read():
    async def scenario():
        entered, release = asyncio.Event(), asyncio.Event()
        async def downstream(scope, receive, send):
            entered.set(); await release.wait()
        middleware = documents.DocumentAdmission(downstream)
        scope = {"type": "http", "method": "POST", "path": "/api/v1/documents/convert"}
        async def receive(): return {"type": "http.request", "body": b"", "more_body": False}
        async def send(message): pass
        first = asyncio.create_task(middleware(scope, receive, send))
        await entered.wait()
        messages = []
        async def forbidden(): raise AssertionError("Busy request body must not be read")
        async def capture(message): messages.append(message)
        await middleware(scope, forbidden, capture)
        assert messages[0]["status"] == 429
        release.set(); await first
        assert not middleware.busy
    asyncio.run(scenario())


def test_worker_timeout_cancellation_and_cleanup(monkeypatch, tmp_path):
    async def scenario():
        process = await asyncio.create_subprocess_exec(__import__("sys").executable, "-c", "import time; time.sleep(60)",
                    stdin=asyncio.subprocess.PIPE, stdout=asyncio.subprocess.PIPE)
        async def spawn(*args, **kwargs): return process
        monkeypatch.setattr(documents.asyncio, "create_subprocess_exec", spawn)
        monkeypatch.setattr(documents, "PARSE_TIMEOUT", 0.05)
        with pytest.raises(ToolFailure) as error:
            await documents.run_worker(tmp_path / "unused", "txt", "a.pdf")
        assert error.value.code == "DOCUMENT_TIMEOUT" and process.returncode is not None
    asyncio.run(scenario())


def test_worker_cancel_terminates_process(monkeypatch, tmp_path):
    async def scenario():
        process = await asyncio.create_subprocess_exec(sys.executable, "-c", "import time; time.sleep(60)",
                    stdin=asyncio.subprocess.PIPE, stdout=asyncio.subprocess.PIPE)
        async def spawn(*args, **kwargs): return process
        monkeypatch.setattr(documents.asyncio, "create_subprocess_exec", spawn)
        task = asyncio.create_task(documents.run_worker(tmp_path / "unused", "txt", "a.pdf"))
        await asyncio.sleep(0.05)
        task.cancel()
        with pytest.raises(asyncio.CancelledError): await task
        assert process.returncode is not None
    asyncio.run(scenario())


def test_worker_hard_memory_limit():
    code = "from app.document_worker import set_limits; set_limits(); " \
           "\ntry: data = bytearray(300 * 1024 * 1024)\nexcept MemoryError: print('LIMIT_ENFORCED')"
    result = subprocess.run([sys.executable, "-c", code], capture_output=True, timeout=8)
    assert result.returncode == 0 and b"LIMIT_ENFORCED" in result.stdout, result.stderr


@pytest.mark.skipif(not sys.platform.startswith("linux"), reason="Verify Linux SIGXCPU in container")
def test_worker_hard_cpu_limit():
    result = subprocess.run([sys.executable, "-c", "from app.document_worker import set_limits; set_limits()\nwhile True: pass"],
                            capture_output=True, timeout=16)
    import signal
    assert result.returncode in {-signal.SIGXCPU, -signal.SIGKILL}


def test_document_download_rejects_private_network():
    response = client.post("/api/v1/documents/convert", data={"file_url": "http://127.0.0.1/a.pdf"})
    assert response.status_code == 400 and response.json()["code"] == "INVALID_INPUT"
